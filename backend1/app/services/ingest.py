"""LangChain load → split → embed → Qdrant + Postgres metadata."""

from __future__ import annotations

import logging
from pathlib import Path

from langchain_community.document_loaders import Docx2txtLoader, PyPDFLoader, TextLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models.entities import ChunkMeta, Document, IngestStatus
from app.services.embeddings import get_embeddings
from app.services.qdrant_store import delete_by_document, upsert_chunks

logger = logging.getLogger(__name__)

# Prefer section / bullet boundaries for KB docs (e.g. 了解山东师范大学)
_SEPARATORS = [
    "\n【",
    "\n## ",
    "\n### ",
    "\n##### ",
    "\n\n",
    "\n- ",
    "\n",
    "。",
    "；",
    " ",
    "",
]


def _load_docs(path: Path):
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return PyPDFLoader(str(path)).load()
    if suffix in {".docx", ".doc"}:
        try:
            return Docx2txtLoader(str(path)).load()
        except Exception:
            from docx import Document as DocxDocument
            from langchain_core.documents import Document as LCDocument

            doc = DocxDocument(str(path))
            text = "\n".join(p.text for p in doc.paragraphs)
            return [LCDocument(page_content=text, metadata={"source": str(path)})]
    return TextLoader(str(path), encoding="utf-8").load()


def _build_splitter():
    settings = get_settings()
    return RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=_SEPARATORS,
        keep_separator=True,
    )


def process_document(db: Session, document_id: str) -> None:
    settings = get_settings()
    doc = db.get(Document, document_id)
    if not doc:
        logger.error("document %s not found", document_id)
        return

    doc.status = IngestStatus.processing
    db.commit()

    try:
        path = Path(doc.storage_path)
        loaded = _load_docs(path)
        splitter = _build_splitter()
        chunks = splitter.split_documents(loaded)
        texts = [c.page_content.strip() for c in chunks if c.page_content.strip()]
        if not texts:
            raise ValueError("no text extracted from document")

        # replace prior vectors for this document (re-ingest / finer split)
        try:
            delete_by_document(doc.id, doc.tenant_id)
        except Exception as exc:  # noqa: BLE001
            logger.warning("qdrant delete skipped: %s", exc)

        embeddings = get_embeddings()
        vectors = embeddings.embed_documents(texts)
        point_ids = upsert_chunks(
            vectors=vectors,
            texts=texts,
            tenant_id=doc.tenant_id,
            document_id=doc.id,
            doc_type=doc.doc_type.value,
            filename=doc.filename,
        )

        db.query(ChunkMeta).filter(ChunkMeta.document_id == doc.id).delete()
        for i, (pid, text) in enumerate(zip(point_ids, texts, strict=True)):
            db.add(
                ChunkMeta(
                    tenant_id=doc.tenant_id,
                    document_id=doc.id,
                    chunk_index=i,
                    qdrant_point_id=pid,
                    content_preview=text[:500],
                )
            )
        doc.chunk_count = len(point_ids)
        doc.status = IngestStatus.succeeded
        doc.error_message = None
        db.commit()
        logger.info("ingest succeeded document=%s chunks=%s", doc.id, len(point_ids))
    except Exception as exc:  # noqa: BLE001
        logger.exception("ingest failed document=%s", document_id)
        doc.status = IngestStatus.failed
        doc.error_message = str(exc)[:2000]
        db.commit()
